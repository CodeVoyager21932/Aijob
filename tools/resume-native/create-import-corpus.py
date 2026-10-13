"""Independent synthetic inputs: browser print / python-docx, never the native renderer.

Every third case is held out from rule tuning, covering all four formats. These are not real user resumes
or a substitute for later real-world validation. Expected labels precede parsing.
"""
from pathlib import Path
import json, html, hashlib
from docx import Document
from docx.shared import Pt
from playwright.sync_api import sync_playwright

root = Path(__file__).resolve().parents[2]
out = root / '.data/resume-native-k2-20260912/independent-corpus'
out.mkdir(parents=True, exist_ok=True)
headings = [
    ('教育经历', '实习经历', '项目经历', '技能'),
    ('教育背景', '工作经历', '项目经验', '专业技能'),
    ('学习经历', '实习与工作经历', '项目实践', '技能与工具'),
    ('教育经历', '实习经历', '科研经历', '语言能力'),
    ('教育背景', '校园实践', '项目经历', '证书'),
]
manifest = []
with sync_playwright() as p:
    browser = p.chromium.launch(channel='chrome', headless=True)
    page = browser.new_page()
    for index in range(30):
        kind = 'pdf-single' if index < 10 else 'pdf-double' if index < 18 else 'docx' if index < 26 else 'txt'
        variant = index % len(headings)
        h = headings[variant]
        name = f'测试同学{index+1}'
        tag = '甲乙丙丁戊己庚辛壬癸子丑寅卯辰巳午未申酉戌亥春夏秋冬东西南北'[index]
        school = f'合成{tag}大学'
        company = f'示例{tag}研究院' if variant == 4 else f'样本{tag}科技有限公司'
        project = f'校园{tag}服务改进项目'
        education = [f'{school}  信息管理  本科', '2023.09 - 2027.06', '成绩：GPA 3.66/4，主修统计学与计算机基础。']
        experience = [f'{company}  产品实习生', '2026.03 - 2026.06', '负责访谈问题整理，完成 18 次用户访谈。', '与设计同学合作完成原型，参与 2 轮可用性测试。']
        projects = [project, '2025.10 - 2026.02', '角色：项目负责人', '整理问卷与使用日志，输出服务改进报告；尚未上线，无营收数据。']
        extra = ['英语 CET-6 520 分'] if variant == 3 else ['大学英语六级  2025.06'] if variant == 4 else ['SQL：使用查询完成留存统计。', 'Python：清理 1200 条合成记录。']
        if index % 3 == 1:
            education[0] = f'{school} | 本科 | 信息管理 | 2023.09—2027.06'
            education.pop(1)
        if index % 3 == 2:
            experience[0] = company
            experience.insert(1, '职位：产品实习生')
        # Distinct difficult structures are declared, not silently excluded.
        if index in (8, 19, 27):
            education += [f'第二合成{tag}大学  管理科学  硕士', '2027.09 - 至今', '研究方向：人机协同与用户研究。']
        if index in (9, 20, 28):
            experience += [f'第二样本{tag}企业  运营实习生', '2025.07 - 2025.09', '整理活动资料，未负责销售。']
        groups = [(h[0], education), (h[1], experience), (h[2], projects), (h[3], extra)]
        header = [name, f'电话：1380000{index:04d}  邮箱：student{index+1}@example.com', '求职意向：产品实习生']
        ext = 'pdf' if kind.startswith('pdf') else kind
        file = out / f'case-{index+1:02d}.{ext}'
        if kind == 'docx':
            document = Document()
            document.styles['Normal'].font.size = Pt(10)
            for line in header: document.add_paragraph(line)
            for title, lines in groups:
                document.add_heading(title, level=1)
                if index % 2:
                    table = document.add_table(rows=1, cols=2)
                    table.cell(0, 0).text = lines[0]
                    table.cell(0, 1).text = lines[1] if len(lines) > 1 else ''
                    for line in lines[2:]: document.add_paragraph(line)
                else:
                    for line in lines: document.add_paragraph(line)
            document.save(file)
        elif kind == 'txt':
            file.write_text('\n'.join(header + [line for title, lines in groups for line in [title, *lines]]), encoding='utf-8')
        else:
            def section(group):
                title, lines = group
                return '<section><h2>'+html.escape(title)+'</h2>'+''.join('<p>'+html.escape(line)+'</p>' for line in lines)+'</section>'
            body = ''.join(section(group) for group in groups)
            if kind == 'pdf-double':
                body = '<div class="columns"><aside>'+section(groups[-1])+'</aside><main>'+''.join(section(group) for group in groups[:-1])+'</main></div>'
            content = '<html lang="zh-CN"><meta charset="utf-8"><style>body{font:10pt "Microsoft YaHei",sans-serif;color:#222}h1{font-size:18pt}h2{font-size:12pt;border-bottom:1px solid #444;margin-top:18pt}p{line-height:1.5;margin:5pt 0}.columns{display:flex;gap:32pt}aside{width:26%}main{flex:1}section{break-inside:avoid}</style><h1>'+html.escape(header[0])+'</h1>'+''.join('<p>'+html.escape(line)+'</p>' for line in header[1:])+body+'</html>'
            page.set_content(content)
            page.pdf(path=str(file), format='A4', margin={side:'16mm' for side in ['top','bottom','left','right']}, print_background=True)
        expected = {'education': [school], 'experience' if variant != 4 else 'volunteer': [company], 'research' if variant == 3 else 'projects':[project]}
        if index in (8,19,27): expected['education'].append(f'第二合成{tag}大学')
        if index in (9,20,28): expected['experience' if variant != 4 else 'volunteer'].append(f'第二样本{tag}企业')
        manifest.append({'id':index+1, 'file':file.name, 'kind':kind, 'split':'holdout' if index%3==2 else 'development', 'expected':expected, 'sha256':hashlib.sha256(file.read_bytes()).hexdigest(), 'origin':'independent synthetic fixture, not a real resume'})
    browser.close()
(out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'{len(manifest)} independently rendered synthetic inputs; 10 frozen holdouts')
