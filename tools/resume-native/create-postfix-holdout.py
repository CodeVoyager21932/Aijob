"""Seal a second independent synthetic corpus before testing the fixed parser.

The first holdout and its failure remain unchanged. This script only generates
new files using ReportLab/python-docx/plain text, never the Aijob renderer.
"""
from pathlib import Path
import hashlib
import json

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / ".data/resume-native-import-acceptance-20260923/postfix-holdout"
if OUT.exists():
    raise SystemExit(f"Holdout already sealed: {OUT}; refusing regeneration")
OUT.mkdir(parents=True)
pdfmetrics.registerFont(TTFont("PostfixCN", r"C:\Windows\Fonts\msyh.ttc", subfontIndex=0))
W, H = A4

CASES = [
    dict(id="P01", kind="pdf-single", name="宋嘉宜", phone="13800002001", email="songjy@example.com",
         sections=[("教育背景", ["南川大学  法学  本科  2022.09—2026.06"]),
                   ("实习经历", ["云岚律师事务所  法务实习生  2025.04—2025.07", "• 核对合同条款并整理法律检索记录。"]),
                   ("项目经历", ["模拟谈判训练  项目成员  2024.08—2024.12", "• 编写案例分析备忘录。"])],
         expected=dict(education=[dict(school="南川大学", degree="本科")],
                       experience=[dict(company="云岚律师事务所", position="法务实习生", evidence="合同条款")],
                       projects=[dict(name="模拟谈判训练", evidence="案例分析")])) ,
    dict(id="P02", kind="pdf-single", name="唐景初", phone="13800002002", email="tangjc@example.com",
         sections=[("教育经历", ["海岳医科大学  统计学  硕士  2023.09—2026.06"]),
                   ("工作经历", ["东陵医院  数据分析实习生  2025.02—2025.06", "• 清理门诊问卷并形成数据字典。"]),
                   ("技能与工具", ["R：完成描述统计。", "Excel：核对原始记录。"])],
         expected=dict(education=[dict(school="海岳医科大学", degree="硕士")],
                       experience=[dict(company="东陵医院", position="数据分析实习生", evidence="门诊问卷")],
                       skills=[dict(name="R"), dict(name="Excel")])) ,
    dict(id="P03", kind="pdf-single", name="乔月宁", phone="13800002003", email="qiaoyn@example.com",
         sections=[("教育背景", ["明珠大学  汉语言文学  本科  2022.09—2026.06"]),
                   ("实习经验", ["松涛出版社  编辑实习生  2025.06—2025.09", "• 校对书稿并记录引文问题。"]),
                   ("语言能力", ["英语：可阅读英文资料。"])],
         expected=dict(education=[dict(school="明珠大学", degree="本科")],
                       experience=[dict(company="松涛出版社", position="编辑实习生", evidence="书稿")],
                       languages=[dict(language="英语")])) ,
    dict(id="P04", kind="pdf-double", name="何致远", phone="13800002004", email="hezy@example.com",
         left=[("教育背景", ["江湾大学  本科  2022.09—2026.06", "社会学专业。"]),
               ("专业技能", ["Excel：整理志愿服务记录。"])],
         right=[("实习经历", ["青石基金会  项目助理  2025.03—2025.07", "• 组织社区访谈并整理需求清单。"]),
                ("校园实践", ["社区读书角  志愿者  2024.04—2024.11", "• 协调活动场地。"])],
         expected=dict(education=[dict(school="江湾大学", degree="本科")],
                       skills=[dict(name="Excel")],
                       experience=[dict(company="青石基金会", position="项目助理", evidence="社区访谈")],
                       volunteer=[dict(organization="社区读书角", evidence="活动场地")])) ,
    dict(id="P05", kind="pdf-double", name="陆予清", phone="13800002005", email="luyq@example.com",
         left=[("教育经历", ["北湖大学  硕士  2023.09—2026.06", "统计学专业。"]),
               ("专业技能", ["Python：处理实验数据。"])],
         right=[("科研经历", ["交通行为研究  研究助理  2024.05—2025.02", "• 完成问卷编码与模型检验。"]),
                ("发表论文", ["论文题目：通勤路径选择研究", "作者：陆予清、王某"])],
         expected=dict(education=[dict(school="北湖大学", degree="硕士")],
                       skills=[dict(name="Python")],
                       research=[dict(name="交通行为研究", evidence="问卷编码")],
                       publications=[dict(title="通勤路径选择研究")])) ,
    dict(id="P06", kind="pdf-double", name="林以安", phone="13800002006", email="linya@example.com",
         left=[("教育背景", ["西山财经大学  本科  2022.09—2026.06", "金融学专业。"]),
               ("资格证书", ["证券从业资格证书  2025.04"])],
         right=[("实习经历", ["北辰银行  财务实习生  2025.06—2025.08", "• 核对报销单据并登记异常。"]),
                ("项目经历", ["校园财务分析  项目成员  2024.03—2024.09", "• 汇总公开财报指标。"])],
         expected=dict(education=[dict(school="西山财经大学", degree="本科")],
                       certifications=[dict(title="证券从业资格证书")],
                       experience=[dict(company="北辰银行", position="财务实习生", evidence="报销单据")],
                       projects=[dict(name="校园财务分析", evidence="财报指标")])) ,
    dict(id="P07", kind="docx", name="陈若思", phone="13800002007", email="chenrs@example.com",
         sections=[("教育经历", ["东桥大学  会计学  本科  2022.09—2026.06"]),
                   ("实习经历", ["启川会计师事务所  审计实习生  2025.01—2025.04", "• 整理往来函证和审计底稿。"]),
                   ("证书", ["大学英语六级  2024.12"])],
         expected=dict(education=[dict(school="东桥大学", degree="本科")],
                       experience=[dict(company="启川会计师事务所", position="审计实习生", evidence="审计底稿")],
                       certifications=[dict(title="大学英语六级")])) ,
    dict(id="P08", kind="docx", name="赵星野", phone="13800002008", email="zhaoxy@example.com",
         sections=[("教育背景", ["远海师范大学  教育学  本科  2022.09—2026.06"]),
                   ("工作经历", ["青芽学校  教学助理  2025.02—2025.06", "• 整理课堂反馈并协助备课。"]),
                   ("获奖经历", ["校级教学创新奖  2025.05"])],
         expected=dict(education=[dict(school="远海师范大学", degree="本科")],
                       experience=[dict(company="青芽学校", position="教学助理", evidence="课堂反馈")],
                       awards=[dict(title="校级教学创新奖")])) ,
    dict(id="P09", kind="text", name="顾晓川", phone="13800002009", email="guxc@example.com",
         sections=[("教育背景", ["岭北大学 法学 本科 2022.09—2026.06"]),
                   ("实习经历", ["万象律所 法务实习生 2025.05—2025.08", "• 撰写案例检索摘要。"]),
                   ("项目实践", ["法律援助项目  项目成员 2024.03—2024.12", "• 整理咨询问题。"])],
         expected=dict(education=[dict(school="岭北大学", degree="本科")],
                       experience=[dict(company="万象律所", position="法务实习生", evidence="案例检索")],
                       projects=[dict(name="法律援助项目", evidence="咨询问题")])) ,
    dict(id="P10", kind="text", name="魏知遥", phone="13800002010", email="weizy@example.com",
         sections=[("教育经历", ["南峰大学 本科 信息管理 2022.09—2026.06"]),
                   ("工作经历", ["2025.06—2025.09 数据分析实习生", "云峰医院", "• 归纳门诊服务意见。"]),
                   ("专业技能", ["SQL：核对调查记录。"])],
         expected=dict(education=[dict(school="南峰大学", degree="本科")],
                       experience=[dict(company="云峰医院", position="数据分析实习生", evidence="门诊服务")],
                       skills=[dict(name="SQL")])) ,
]

def lines(text, width, size):
    row = ""
    for char in text:
        if row and pdfmetrics.stringWidth(row + char, "PostfixCN", size) > width:
            yield row
            row = char
        else:
            row += char
    if row:
        yield row

def draw_groups(pdf, groups, x, y, width):
    for title, values in groups:
        pdf.setFont("PostfixCN", 11)
        pdf.drawString(x, y, title)
        y -= 23
        for value in values:
            pdf.setFont("PostfixCN", 9)
            for row in lines(value, width, 9):
                pdf.drawString(x, y, row)
                y -= 16
        y -= 17

manifest = []
for case in CASES:
    suffix = "pdf" if case["kind"].startswith("pdf") else "docx" if case["kind"] == "docx" else "txt"
    path = OUT / f"{case['id']}.{suffix}"
    if suffix == "pdf":
        pdf = canvas.Canvas(str(path), pagesize=A4)
        pdf.setFont("PostfixCN", 16)
        pdf.drawString(42, H - 46, case["name"])
        pdf.setFont("PostfixCN", 9)
        pdf.drawString(42, H - 65, f"电话：{case['phone']}  邮箱：{case['email']}")
        if case["kind"] == "pdf-double":
            draw_groups(pdf, case["left"], 42, H - 116, 195)
            draw_groups(pdf, case["right"], 280, H - 116, 270)
        else:
            draw_groups(pdf, case["sections"], 42, H - 102, W - 84)
        pdf.save()
    elif suffix == "docx":
        doc = Document()
        doc.styles["Normal"].font.size = Pt(10)
        doc.add_heading(case["name"], level=1)
        doc.add_paragraph(f"电话：{case['phone']}  邮箱：{case['email']}")
        for title, values in case["sections"]:
            doc.add_heading(title, level=2)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = values[0]
            for value in values[1:]:
                doc.add_paragraph(value)
        doc.save(path)
    else:
        rows = [case["name"], f"电话：{case['phone']}  邮箱：{case['email']}"]
        for title, values in case["sections"]:
            rows.extend([title, *values])
        path.write_text("\n".join(rows), encoding="utf-8")
    manifest.append({"id": case["id"], "file": path.name, "kind": case["kind"],
                     "split": "postfix-holdout", "expected": {
                         "basics": {key: case[key] for key in ("name", "phone", "email")},
                         **case["expected"]}, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                     "origin": "synthetic; independently rendered with ReportLab, python-docx or plain text"})
manifest_bytes = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
(OUT / "manifest.json").write_bytes(manifest_bytes)
(OUT / "seal.json").write_text(json.dumps({
    "count": len(manifest),
    "manifestSha256": hashlib.sha256(manifest_bytes).hexdigest(),
    "formats": {kind: sum(item["kind"] == kind for item in manifest)
                for kind in sorted({item["kind"] for item in manifest})}}, indent=2), encoding="utf-8")
print((OUT / "seal.json").read_text(encoding="utf-8"))
