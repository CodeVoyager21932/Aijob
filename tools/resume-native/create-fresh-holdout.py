"""Freeze ten new, synthetic import inputs from non-Aijob renderers.

These fixtures must never be used to choose parser rules. Run once, score only after
rules are frozen, and record failures rather than editing generated documents.
"""
from __future__ import annotations

from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.shared import Cm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / ".data/resume-native-import-acceptance-20260923/fresh-holdout"
if OUT.exists():
    raise SystemExit(f"Holdout already sealed: {OUT}; refusing regeneration")
OUT.mkdir(parents=True)
FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
if not FONT.exists():
    raise SystemExit("The independent ReportLab renderer requires Microsoft YaHei on this host")
pdfmetrics.registerFont(TTFont("HoldoutCN", str(FONT), subfontIndex=0))
W, H = A4

# The expected categories and fields below are written before any Aijob parser run.
CASES = [
    dict(id="H01", kind="pdf-single", name="林若舟", phone="13800001001", email="linrz@example.com",
         sections=[("教育背景", ["云岭大学  信息管理  本科  2022.09—2026.06", "主修统计学与信息系统。"]),
                   ("实习经历", ["远舟科技有限公司  产品实习生  2025.06—2025.09", "• 整理用户访谈记录，输出需求文档。"]),
                   ("项目经历", ["校园地图改进项目  项目负责人  2024.10—2025.03", "• 完成原型与可用性测试。"]),
                   ("专业技能", ["SQL：整理问卷与行为数据。"])],
         expected=dict(education=[dict(school="云岭大学", degree="本科", period="2022.09—2026.06")],
                       experience=[dict(company="远舟科技有限公司", position="产品实习生", evidence="用户访谈记录")],
                       projects=[dict(name="校园地图改进项目", evidence="可用性测试")], skills=[dict(name="SQL")])),
    dict(id="H02", kind="pdf-single", name="周念一", phone="13800001002", email="zhouny@example.com",
         sections=[("一、教育经历", ["海川财经大学  会计学  本科  2021.09—2025.06", "北方商学院  财务管理  硕士  2025.09—2027.06"]),
                   ("二、实习与工作经历", ["南星咨询有限公司  审计实习生  2024.07—2024.10", "• 核对采购凭证并记录异常项。", "衡岳科技有限公司  财务实习生  2025.07—2025.09", "• 整理月度费用明细。"]),
                   ("三、资格证书", ["初级会计资格证书  2025.05"])],
         expected=dict(education=[dict(school="海川财经大学", degree="本科"),dict(school="北方商学院", degree="硕士")],
                       experience=[dict(company="南星咨询有限公司", evidence="采购凭证"),dict(company="衡岳科技有限公司", evidence="月度费用")],
                       certifications=[dict(title="初级会计资格证书")])),
    dict(id="H03", kind="pdf-single", name="顾清禾", phone="13800001003", email="guqh@example.com",
         sections=[("教育背景", ["岭南医科大学  生物医学  硕士  2023.09—2026.06", "研究方向：公共健康数据。"]),
                   ("科研经历", ["社区健康访谈研究  研究助理  2024.03—2025.01", "• 完成访谈编码和质性分析。"]),
                   ("论文成果", ["论文题目：社区健康服务需求研究", "作者：顾清禾、李某", "发表状态：在投"]),
                   ("语言能力", ["英语：能够阅读专业文献。"])],
         expected=dict(education=[dict(school="岭南医科大学", degree="硕士")],
                       research=[dict(name="社区健康访谈研究", evidence="质性分析")],
                       publications=[dict(title="社区健康服务需求研究")], languages=[dict(language="英语")])),
    dict(id="H04", kind="pdf-double", name="陈以北", phone="13800001004", email="chenyb@example.com",
         left=[("教育经历", ["江城理工大学  本科", "计算机科学与技术  2022.09—2026.06"]),
               ("技能与工具", ["Python：编写清洗脚本。", "Linux：部署测试服务。"])],
         right=[("项目经历", ["校园检索系统  开发负责人  2024.09—2025.02", "• 建立文档索引和查询接口。"]),
                ("实习经历", ["启航网络有限公司  后端开发实习生  2025.06—2025.08", "• 修复慢查询并补充自动化测试。"])],
         expected=dict(education=[dict(school="江城理工大学", degree="本科")],
                       skills=[dict(name="Python"),dict(name="Linux")],
                       projects=[dict(name="校园检索系统", evidence="文档索引")],
                       experience=[dict(company="启航网络有限公司", evidence="慢查询")])),
    dict(id="H05", kind="pdf-double", name="许书瑶", phone="13800001005", email="xusy@example.com",
         left=[("教育背景", ["东海艺术学院  本科", "视觉传达设计  2022.09—2026.06"]),
               ("专业技能", ["Figma：绘制交互原型。", "Photoshop：处理视觉素材。"])],
         right=[("项目实践", ["新生办事指南设计  视觉设计  2024.04—2024.10", "• 绘制信息架构并交付页面规范。"]),
                ("实习经历", ["知行设计工作室  设计实习生  2025.07—2025.09", "• 参与海报和活动页面设计。"])],
         expected=dict(education=[dict(school="东海艺术学院", degree="本科")],
                       skills=[dict(name="Figma"),dict(name="Photoshop")],
                       projects=[dict(name="新生办事指南设计", evidence="信息架构")],
                       experience=[dict(company="知行设计工作室", evidence="活动页面")])),
    dict(id="H06", kind="pdf-double", name="陆见星", phone="13800001006", email="lujx@example.com",
         left=[("教育情况", ["青禾大学  本科", "新闻传播学  2022.09—2026.06"]),
               ("技能清单", ["Excel：整理活动反馈。", "剪映：完成短视频剪辑。"])],
         right=[("实习经验", ["星云文化有限公司  内容运营实习生  2025.03—2025.07", "• 制作内容排期并复盘阅读数据。"]),
                ("校园活动", ["新生读书会  组织者  2024.09—2024.12", "• 协调志愿者与场地安排。"])],
         expected=dict(education=[dict(school="青禾大学", degree="本科")],
                       skills=[dict(name="Excel"),dict(name="剪映")],
                       experience=[dict(company="星云文化有限公司", evidence="内容排期")],
                       volunteer=[dict(organization="新生读书会", evidence="志愿者")])),
    dict(id="H07", kind="docx", name="苏遥", phone="13800001007", email="suyao@example.com",
         sections=[("教育经历", ["西洲大学  本科  市场营销  2022.09—2026.06"]),
                   ("实习经历", ["晨星品牌有限公司  市场实习生  2025.06—2025.09", "• 策划校园问卷并整理反馈。"]),
                   ("项目经历", ["社区品牌调研  项目成员  2024.11—2025.02", "• 访谈店主并制作调研报告。"]),
                   ("技能与工具", ["Excel：制作数据透视表。"])],
         expected=dict(education=[dict(school="西洲大学", degree="本科")],
                       experience=[dict(company="晨星品牌有限公司", evidence="校园问卷")],
                       projects=[dict(name="社区品牌调研", evidence="店主")],skills=[dict(name="Excel")])),
    dict(id="H08", kind="docx", name="魏予安", phone="13800001008", email="weiya@example.com",
         sections=[("教育背景", ["北辰大学  硕士  统计学  2023.09—2026.06"]),
                   ("研究经历", ["城市出行数据研究  研究助理  2024.01—2025.01", "• 清理调查数据并撰写方法说明。"]),
                   ("发表论文", ["论文题目：公共交通调查方法比较", "作者：魏予安、赵某", "发表状态：已接收"]),
                   ("获奖经历", ["校级研究奖  2025.06"])],
         expected=dict(education=[dict(school="北辰大学", degree="硕士")],
                       research=[dict(name="城市出行数据研究", evidence="调查数据")],
                       publications=[dict(title="公共交通调查方法比较")],awards=[dict(title="校级研究奖")])),
    dict(id="H09", kind="text", name="孟知夏", phone="13800001009", email="mengzx@example.com",
         sections=[("教育经历", ["锦城大学 本科 法学 2022.09—2026.06"]),
                   ("实习经历", ["清源律所 法务实习生 2025.07—2025.09", "• 整理合同版本，标注争议条款。"]),
                   ("校园实践", ["模拟法庭  队员  2024.03—2024.11", "• 检索案例并整理答辩材料。"])],
         expected=dict(education=[dict(school="锦城大学",degree="本科")],
                       experience=[dict(company="清源律所",evidence="合同版本")],
                       volunteer=[dict(organization="模拟法庭",evidence="答辩材料")])),
    dict(id="H10", kind="text", name="方启明", phone="13800001010", email="fangqm@example.com",
         sections=[("教育背景", ["明川大学 本科 自动化 2022.09—2026.06"]),
                   ("项目成果", ["智能温室监测平台  项目成员  2024.06—2025.04", "• 搭建传感器数据采集流程。"]),
                   ("实习经验", ["东岭制造有限公司  工程实习生  2025.06—2025.08", "• 参与设备测试并记录缺陷。"]),
                   ("证书", ["大学英语六级  2024.12"])],
         expected=dict(education=[dict(school="明川大学",degree="本科")],
                       projects=[dict(name="智能温室监测平台",evidence="传感器数据")],
                       experience=[dict(company="东岭制造有限公司",evidence="设备测试")],
                       certifications=[dict(title="大学英语六级")])),
]

def split_lines(text: str, max_width: float, size: float) -> list[str]:
    result, line = [], ""
    for char in text:
        if line and pdfmetrics.stringWidth(line + char, "HoldoutCN", size) > max_width:
            result.append(line)
            line = char
        else:
            line += char
    if line:
        result.append(line)
    return result


def pdf_rows(c: canvas.Canvas, groups, x: float, y: float, width: float, page_safe: bool) -> float:
    for title, values in groups:
        if page_safe and y < 92:
            c.showPage()
            y = H - 55
        c.setFont("HoldoutCN", 11)
        c.drawString(x, y, title)
        y -= 24
        for index, value in enumerate(values):
            size = 9
            c.setFont("HoldoutCN", size)
            for line in split_lines(value, width, size):
                if page_safe and y < 58:
                    c.showPage()
                    y = H - 55
                    c.setFont("HoldoutCN", size)
                c.drawString(x, y, line)
                y -= 17
            y -= 3 if index == 0 else 0
        y -= 12
    return y


def write_pdf(path: Path, case: dict) -> None:
    c = canvas.Canvas(str(path), pagesize=A4, pageCompression=1)
    c.setFont("HoldoutCN", 16)
    c.drawString(42, H - 46, case["name"])
    c.setFont("HoldoutCN", 9)
    c.drawString(42, H - 65, f"电话：{case['phone']}  邮箱：{case['email']}")
    if case["kind"] == "pdf-double":
        pdf_rows(c, case["left"], 42, H - 122, 194, False)
        pdf_rows(c, case["right"], 280, H - 122, 270, False)
    else:
        pdf_rows(c, case["sections"], 42, H - 103, W - 84, True)
    c.save()


def write_docx(path: Path, case: dict) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    for index, value in enumerate([case["name"], f"电话：{case['phone']}  邮箱：{case['email']}"]):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(value)
        run.bold = index == 0
        run.font.size = Pt(15 if index == 0 else 9)
    for title, values in case["sections"]:
        document.add_heading(title, level=2)
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = values[0]
        for value in values[1:]:
            paragraph = document.add_paragraph(value)
            paragraph.style = "List Bullet" if value.startswith("•") else "Normal"
    document.save(path)

manifest=[]
for case in CASES:
    suffix = "pdf" if case["kind"].startswith("pdf") else "docx" if case["kind"] == "docx" else "txt"
    path = OUT / f"{case['id']}.{suffix}"
    if suffix == "pdf":
        write_pdf(path, case)
    elif suffix == "docx":
        write_docx(path, case)
    else:
        lines = [case["name"], f"电话：{case['phone']}  邮箱：{case['email']}"]
        for title, values in case["sections"]:
            lines.extend([title, *values])
        path.write_text("\n".join(lines), encoding="utf-8")
    manifest.append({"id": case["id"], "file": path.name, "kind": case["kind"],
                     "split": "fresh-holdout", "expected": {"basics": {key: case[key] for key in ("name","phone","email")}, **case["expected"]},
                     "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                     "origin": "synthetic; independently rendered with ReportLab or python-docx or plain text"})
manifest_bytes = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
(OUT / "manifest.json").write_bytes(manifest_bytes)
(OUT / "seal.json").write_text(json.dumps({"count": len(manifest),
    "manifestSha256": hashlib.sha256(manifest_bytes).hexdigest(),
    "formats": {kind: sum(x["kind"] == kind for x in manifest) for kind in sorted({x["kind"] for x in manifest})}}, indent=2), encoding="utf-8")
print(OUT)
print((OUT / "seal.json").read_text(encoding="utf-8"))
